"""Écriture des sorties : releves.csv, releves.xlsx, releves.geojson, rapport_import.txt."""
from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .consolidation import (COLONNES_ANOMALIES, COLONNES_RELEVES, COLONNES_VISITES, STATUT_CORBEILLE,
                            STATUT_DOUBLON, TYPES_CONNUS, Resultat)
from .noms import chemin_os

COLONNES_LIENS = ("fichier_photo", "fichier_audio", "fichier_joint", "fichier_source")
_INTERDITS_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _texte_csv(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    return str(v)


def ecrire_csv(chemin: Path, lignes: list[dict], separateur: str) -> None:
    # UTF-8 avec BOM : Excel sous Windows reconnaît ainsi les accents.
    with open(chemin_os(chemin), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=separateur, quoting=csv.QUOTE_MINIMAL)
        w.writerow(COLONNES_RELEVES)
        for l in lignes:
            w.writerow([_texte_csv(l[c]) for c in COLONNES_RELEVES])


def _feuille(wb: Workbook, titre: str, colonnes: list[str], lignes: list[dict], premiere: bool = False) -> None:
    ws = wb.active if premiere else wb.create_sheet()
    ws.title = titre
    ws.append(colonnes)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    lien = Font(color="0563C1", underline="single")
    for i, l in enumerate(lignes, start=2):
        for j, c in enumerate(colonnes, start=1):
            v = l.get(c)
            if isinstance(v, str):
                v = _INTERDITS_XML.sub("", v)
            cell = ws.cell(row=i, column=j)
            cell.value = v if v != "" else None
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"            # texte, jamais une formule
            if isinstance(v, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if c in COLONNES_LIENS[:3] and v:
                cell.hyperlink = v              # lien relatif au classeur : photos/…, audio/…, fichiers/…
                cell.font = lien
    ws.freeze_panes = "A2"
    if lignes:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colonnes))}{len(lignes) + 1}"
    else:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colonnes))}1"
    larges = {"texte": 60, "anomalie": 60, "message": 70, "visite_uid": 24, "fichier_photo": 40,
              "fichier_audio": 40, "fichier_joint": 40, "fichier_source": 40, "date_heure_locale": 20,
              "date_locale": 20}
    for j, c in enumerate(colonnes, start=1):
        ws.column_dimensions[get_column_letter(j)].width = larges.get(c, max(10, len(c) + 2))


def ecrire_xlsx(chemin: Path, res: Resultat, lignes_visites: list[dict]) -> None:
    wb = Workbook()
    _feuille(wb, "releves", COLONNES_RELEVES, res.lignes, premiere=True)
    _feuille(wb, "visites", COLONNES_VISITES, lignes_visites)
    _feuille(wb, "anomalies", COLONNES_ANOMALIES, [a.ligne() for a in res.anomalies])
    wb.save(chemin_os(chemin))


def ecrire_geojson(chemin: Path, lignes: list[dict]) -> int:
    """Un point par entrée positionnée, WGS84, coordonnées [lng, lat]."""
    features = []
    for l in lignes:
        if l["lat"] is None or l["lng"] is None:
            continue
        props = {c: (_texte_csv(l[c]) if isinstance(l[c], datetime) else l[c]) for c in COLONNES_RELEVES}
        features.append({"type": "Feature",
                         "geometry": {"type": "Point", "coordinates": [l["lng"], l["lat"]]},
                         "properties": props})
    fc = {"type": "FeatureCollection", "features": features}   # RFC 7946 : WGS84 implicite
    with open(chemin_os(chemin), "w", encoding="utf-8") as f:
        json.dump(fc, f, ensure_ascii=False, indent=1)
    return len(features)


def ecrire_rapport(chemin: Path, res: Resultat, source: str, sortie: Path, fuseau: str, config_source: str,
                   inclure_corbeille: bool, nb_points: int) -> str:
    r: list[str] = []
    add = r.append
    add("CALEPIN BUREAU — RAPPORT D'IMPORT")
    add("=" * 40)
    add(f"Date de l'import : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    add(f"Source           : {source}")
    add(f"Sortie           : {sortie}")
    add(f"Fuseau           : {fuseau}")
    add(f"Configuration    : {config_source}")
    add(f"Corbeille        : {'incluse' if inclure_corbeille else 'exclue'}")
    add("")
    statuts = Counter(v.statut for v in res.visites)
    add("VISITES")
    add("-------")
    add(f"Digests lus                  : {res.nb_digests_lus}")
    add(f"Visites importées            : {len(res.retenues)}")
    add(f"Doublons écartés             : {statuts.get(STATUT_DOUBLON, 0)}")
    add(f"En corbeille, exclues        : {statuts.get(STATUT_CORBEILLE, 0)}")
    add("")
    for v in res.visites:
        add(f"  [{v.statut}] {v.nom or '(sans nom)'} — {v.uid or '(sans uid)'} — schéma {v.schema or '?'} — "
            f"{len(v.d.get('entries') or [])} entrée(s) — {v.digest.source}")
    add("")
    types = Counter(l["type"] or "(sans type)" for l in res.lignes)
    add("ENTRÉES PAR TYPE (visites importées)")
    add("------------------------------------")
    for t in list(TYPES_CONNUS) + sorted(set(types) - set(TYPES_CONNUS)):
        add(f"  {t:<10} {types.get(t, 0)}")
    add(f"  {'total':<10} {sum(types.values())}")
    add(f"  Points dans le GeoJSON : {nb_points}")
    add("")
    add("FICHIERS EXTRAITS")
    add("-----------------")
    for dossier in ("photos", "audio", "fichiers"):
        add(f"  {dossier + '/':<10} {sum(1 for f in res.fichiers_ecrits if f.startswith(dossier + '/'))}")
    add("")
    if res.avertissements:
        add("AVERTISSEMENTS")
        add("--------------")
        for a in res.avertissements:
            add(f"  ! {a}")
        add("")
    anomalies = res.anomalies
    add(f"ANOMALIES ({len(anomalies)})")
    add("-" * (len(f"ANOMALIES ({len(anomalies)})")))
    for (grav, code), n in sorted(Counter((a.gravite, a.code) for a in anomalies).items(),
                                  key=lambda x: ({"erreur": 0, "attention": 1}.get(x[0][0], 2), x[0][1])):
        add(f"  {grav:<9} {code:<26} {n}")
    add("")
    for a in anomalies:
        ou = " / ".join(x for x in (a.visite_nom or a.visite_uid, a.entree_id and f"entrée {a.entree_id}",
                                    a.type, a.repere and f"« {a.repere} »") if x)
        add(f"  [{a.gravite}] {a.code} — {ou or a.fichier_source} : {a.message}")
    texte = "\n".join(r) + "\n"
    with open(chemin_os(chemin), "w", encoding="utf-8") as f:
        f.write(texte)
    return texte
