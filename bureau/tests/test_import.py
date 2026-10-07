"""Tests sur le jeu synthétique (format tiré du code de l'app)."""
from __future__ import annotations

import base64
import csv
import hashlib
import json
import shutil
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from calepin_bureau.cli import main
from calepin_bureau.config import Config, charger_config
from calepin_bureau.consolidation import theme_de
from calepin_bureau.importer import ErreurImport, importer
from calepin_bureau.noms import assainir, nom_media


def ligne(res, entree_id):
    return next(l for l in res.lignes if l["entree_id"] == entree_id)


def codes(res, entree_id=None, uid=None):
    return {a.code for a in res.anomalies
            if (entree_id is None or a.entree_id == entree_id) and (uid is None or a.visite_uid == uid)}


def empreintes(dossier: Path) -> dict:
    return {p.relative_to(dossier).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(dossier.rglob("*")) if p.is_file()}


# ------------------------------------------------------------------ comptes, schémas, doublons

def test_comptes_par_type(import_jeu):
    res, rapport, _ = import_jeu
    assert Counter(l["type"] for l in res.lignes) == {"photo": 4, "audio": 3, "texte": 5, "point": 2, "fichier": 2}
    assert len(res.retenues) == 2
    assert "photo      4" in rapport and "total      16" in rapport


def test_schemas_1_et_2(import_jeu):
    res, _, _ = import_jeu
    schemas = {v.uid: v.schema for v in res.retenues}
    assert schemas == {"uid-aaaa-1111": 2, "uid-cccc-3333": 1}
    assert "SCHEMA_1" in codes(res, uid="uid-cccc-3333")
    # v1 : audio sans fichier (normal), transcription conservée
    audio_v1 = [l for l in res.lignes if l["visite_uid"] == "uid-cccc-3333" and l["type"] == "audio"][0]
    assert audio_v1["fichier_audio"] == "" and audio_v1["texte"] == "Transcription v1"
    assert audio_v1["source_texte"] == "transcrit" and "PIECE_ILLISIBLE" not in audio_v1["anomalie"]
    assert audio_v1["entree_id"] == ""   # pas d'id en v1 : rien n'est inventé


def test_doublons_garde_le_plus_recent(import_jeu):
    res, _, _ = import_jeu
    groupe = [v for v in res.visites if v.uid == "uid-aaaa-1111"]
    assert len(groupe) == 2
    gardee = next(v for v in groupe if v.statut == "importée")
    assert gardee.d["genereLe"] == "2026-10-05T12:00:00.000Z"
    assert len(gardee.lignes) == 13
    assert {"DOUBLON_GARDE", "DOUBLON_ECARTE"} <= codes(res, uid="uid-aaaa-1111")


def test_corbeille_exclue_puis_incluse(jeu, import_jeu, tmp_path):
    res, _, _ = import_jeu
    assert next(v for v in res.visites if v.uid == "uid-bbbb-2222").statut == "corbeille, exclue"
    assert not [l for l in res.lignes if l["visite_uid"] == "uid-bbbb-2222"]
    res2, _ = importer(jeu["dossier"], tmp_path / "s", Config(), inclure_corbeille=True)
    assert [l["entree_id"] for l in res2.lignes if l["visite_uid"] == "uid-bbbb-2222"] == ["b01"]


# ------------------------------------------------------------------ positions et anomalies

def test_position_nulle(import_jeu):
    res, _, sortie = import_jeu
    l = ligne(res, "e09")
    assert l["lat"] is None and l["lng"] is None and l["position_ok"] == "non"
    assert l["position_motif"] == "refusee"
    assert "POSITION_ABSENTE" in l["anomalie"] and "autorisation refusée" in l["anomalie"]
    ids = {f["properties"]["entree_id"] for f in json.loads((sortie / "releves.geojson").read_text("utf-8"))["features"]}
    assert "e09" not in ids


def test_precision_et_altitude(import_jeu):
    res, _, _ = import_jeu
    l = ligne(res, "e08")
    assert l["precision_m"] == 120 and l["altitude_m"] == 1510.4 and l["position_ok"] == "non"
    assert "PRECISION_FAIBLE" in l["anomalie"]
    assert l["gps_releve_utc"] == "2026-10-05T08:06:59Z"
    assert "PRECISION_INCONNUE" in codes(res, "e13")
    assert ligne(res, "e01")["position_ok"] == "oui"


def test_seuil_precision_configurable(jeu, tmp_path):
    res, _ = importer(jeu["dossier"], tmp_path / "s", Config(precision_max_m=200))
    assert "PRECISION_FAIBLE" not in codes(res, "e08")


def test_piece_illisible(import_jeu):
    res, _, _ = import_jeu
    assert "PIECE_ILLISIBLE" in codes(res, "e03")
    assert ligne(res, "e03")["fichier_photo"] == ""
    assert "PIECE_ILLISIBLE" in codes(res, "n° 3", uid="uid-cccc-3333")    # fichier v1 avec « erreur »


def test_base64_corrompu(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    digest = {"schemaVersion": 2, "uid": "u1", "tag": "t", "date": "2026-10-05T07:00:00Z", "dossier": None,
              "tags": [], "corbeille": None, "appareil": "a", "versionApp": "v", "genereLe": "2026-10-05T07:00:00Z",
              "entries": [{"id": "x", "type": "photo", "horodatage": "2026-10-05T07:00:00Z", "repere": "ERO 1",
                           "position": None, "positionMotif": "delai_depasse", "texte": None,
                           "image": {"mime": "image/jpeg", "base64": base64.b64encode(b"\xff\xd8\xffpas une image").decode()}}]}
    (d / "a.json").write_text(json.dumps(digest), encoding="utf-8")
    res, _ = importer(d, tmp_path / "out", Config())
    l = res.lignes[0]
    assert "PIECE_ILLISIBLE" in l["anomalie"] and "délai GPS dépassé" in l["anomalie"]


def test_photo_sans_exif_gps(import_jeu):
    res, _, _ = import_jeu
    assert "PHOTO_SANS_EXIF_GPS" in codes(res, "e02")
    assert "PHOTO_SANS_EXIF_GPS" not in codes(res, "e01")


def test_sans_repere(import_jeu):
    res, _, _ = import_jeu
    assert "SANS_REPERE" in codes(res, "e11")
    assert ligne(res, "e11")["theme"] == "AUTRE"


def test_audio_silencieux(import_jeu):
    res, _, _ = import_jeu
    if shutil.which("ffmpeg"):
        assert "AUDIO_SILENCIEUX" in codes(res, "e05")
        assert "AUDIO_SILENCIEUX" not in codes(res, "e04")
    else:
        assert any("ffmpeg" in a for a in res.avertissements)


def test_ffmpeg_absent_avertit(jeu, tmp_path):
    res, rapport = importer(jeu["dossier"], tmp_path / "s", Config(ffmpeg=str(tmp_path / "pas-de-ffmpeg.exe")))
    assert any("ffmpeg introuvable" in a for a in res.avertissements)
    assert "AUDIO_SILENCIEUX" not in codes(res)
    assert "ffmpeg introuvable" in rapport


# ------------------------------------------------------------------ textes et thèmes

def test_source_texte(import_jeu):
    res, _, _ = import_jeu
    assert ligne(res, "e06")["source_texte"] == "saisi"
    assert ligne(res, "e07")["source_texte"] == "dictee"
    assert ligne(res, "e05")["source_texte"] == "transcrit"
    assert ligne(res, "e04")["source_texte"] == "" and ligne(res, "e04")["texte"] == ""


@pytest.mark.parametrize("repere,theme", [
    ("FONC 12 borne", "FONC"), ("fonc-12", "FONC"), ("OCC–4 champ", "OCC"), ("HYD", "HYD"),
    ("TRAV 1", "TRAV"), ("ACT: marché", "ACT"), ("XYZ 9", "AUTRE"), ("FONCIER 1", "AUTRE"),
    ("FONC12", "AUTRE"), ("", "AUTRE"), (None, "AUTRE"), ("  riv 5", "RIV"),
])
def test_codes_theme(repere, theme):
    assert theme_de(repere, Config()) == theme


def test_themes_modifiables(tmp_path):
    f = tmp_path / "config.yaml"
    f.write_text("themes: [FONC, PONT]\ntheme_defaut: DIVERS\n", encoding="utf-8")
    c = charger_config(f)
    assert theme_de("pont 3", c) == "PONT" and theme_de("HYD 1", c) == "DIVERS"


def test_config_du_depot_lisible():
    c = charger_config(Path(__file__).parent.parent / "config.yaml")
    assert c.themes == ["FONC", "OCC", "SOC", "ACC", "HYD", "RIV", "TRAV", "ERO", "EMP", "ACT"]
    assert c.precision_max_m == 50 and c.audio_silence_db == -60 and c.fuseau == "Africa/Kigali"


def test_themes_dans_les_lignes(import_jeu):
    res, _, _ = import_jeu
    assert {l["entree_id"]: l["theme"] for l in res.lignes if l["entree_id"] in ("e01", "e02", "e06", "e12")} == \
        {"e01": "FONC", "e02": "HYD", "e06": "OCC", "e12": "AUTRE"}


# ------------------------------------------------------------------ fichiers extraits

def test_photos_extraites_sans_recompression(jeu, import_jeu):
    res, _, sortie = import_jeu
    p1 = sortie / ligne(res, "e01")["fichier_photo"]
    p2 = sortie / ligne(res, "e02")["fichier_photo"]
    assert p1.read_bytes() == jeu["photo_gps"] and p2.read_bytes() == jeu["photo_sans"]
    for p in (sortie / "photos").iterdir():
        with Image.open(p) as img:
            img.load()
            assert img.format == "JPEG"


def test_audio_et_pieces_jointes_d_origine(jeu, import_jeu):
    res, _, sortie = import_jeu
    a = ligne(res, "e04")
    assert a["fichier_audio"].endswith(".webm") and a["duree_audio_s"] == 2
    assert (sortie / a["fichier_audio"]).read_bytes() == jeu["son"]
    f = ligne(res, "e10")["fichier_joint"]
    assert f == "fichiers/Site-Mbandaka-rive-gauche_10_Plan-d-emprise-e.pdf"
    assert (sortie / f).read_bytes() == b"%PDF-1.4 test"


def test_noms_de_fichiers(import_jeu):
    res, _, sortie = import_jeu
    assert ligne(res, "e01")["fichier_photo"] == "photos/Site-Mbandaka-rive-gauche_01_FONC-12-borne.jpg"
    for f in res.fichiers_ecrits:
        nom = f.split("/")[1]
        assert len(nom) <= 120 and nom.isascii() and not set(nom) & set('<>:"/\\|?* ')


def test_assainir():
    assert assainir("Éro: rive / gauche* ?") == "Ero-rive-gauche"
    assert assainir("Cœur d'Ève") == "Coeur-d-Eve"
    assert assainir("CON") == "_CON" and assainir("") == "sans-nom" and assainir("...") == "sans-nom"
    nom = nom_media("v" * 200, "07", "r" * 300, ".jpg", 120)
    assert len(nom) <= 120 and nom.endswith("_07_" + nom.split("_07_")[1]) and nom.endswith(".jpg")


def test_meme_nom_de_visite(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    for uid in ("11111111-a", "22222222-b"):
        digest = {"schemaVersion": 2, "uid": uid, "tag": "Site 1", "date": "2026-10-05T07:00:00Z", "dossier": None,
                  "tags": [], "corbeille": None, "appareil": "a", "versionApp": "v",
                  "genereLe": "2026-10-05T07:00:00Z",
                  "entries": [{"id": "p", "type": "fichier", "horodatage": "2026-10-05T07:00:00Z", "repere": "R",
                               "position": None, "texte": None,
                               "fichier": {"nom": "a.txt", "mime": "text/plain", "base64": "YQ=="}}]}
        (d / f"{uid}.json").write_text(json.dumps(digest), encoding="utf-8")
    res, _ = importer(d, tmp_path / "out", Config())
    noms = [l["fichier_joint"] for l in res.lignes]
    assert len(set(noms)) == 2 and all("Site-1-" in n for n in noms)


# ------------------------------------------------------------------ sorties tabulaires et GeoJSON

def test_geojson_valide_lng_lat(jeu, import_jeu):
    res, _, sortie = import_jeu
    fc = json.loads((sortie / "releves.geojson").read_text("utf-8"))
    assert fc["type"] == "FeatureCollection"
    positionnees = [l for l in res.lignes if l["lat"] is not None]
    assert len(fc["features"]) == len(positionnees) == 13
    sources = {e["id"]: e for e in jeu["visite_a"]["entries"]}
    for f in fc["features"]:
        assert f["type"] == "Feature" and f["geometry"]["type"] == "Point"
        lng, lat = f["geometry"]["coordinates"]
        assert -180 <= lng <= 180 and -90 <= lat <= 90
        assert lng == f["properties"]["lng"] and lat == f["properties"]["lat"]
        src = sources.get(f["properties"]["entree_id"])
        if src:
            assert [lng, lat] == [src["position"]["lng"], src["position"]["lat"]]
    # Kigali / RDC : longitude ~30 E, latitude ~-2 : l'ordre inverse sauterait aux yeux
    assert fc["features"][0]["geometry"]["coordinates"] == [30.06123, -1.95421]


def test_csv(import_jeu):
    res, _, sortie = import_jeu
    brut = (sortie / "releves.csv").read_bytes()
    assert brut.startswith(b"\xef\xbb\xbf")
    lignes = list(csv.DictReader((sortie / "releves.csv").open(encoding="utf-8-sig"), delimiter=";"))
    assert len(lignes) == len(res.lignes)
    assert list(lignes[0])[:22] == [
        "visite_uid", "visite_nom", "appareil", "version_app", "entree_id", "type", "date_heure_locale",
        "date_heure_utc", "theme", "repere", "lat", "lng", "precision_m", "altitude_m", "position_ok", "anomalie",
        "texte", "source_texte", "fichier_photo", "fichier_audio", "duree_audio_s", "fichier_joint"]
    e01 = next(l for l in lignes if l["entree_id"] == "e01")
    assert e01["date_heure_utc"] == "2026-10-05T08:00:02Z"
    assert e01["date_heure_locale"] == "2026-10-05 10:00:02"     # Kigali = UTC+2


def test_xlsx(import_jeu):
    res, _, sortie = import_jeu
    wb = load_workbook(sortie / "releves.xlsx")
    assert wb.sheetnames == ["releves", "visites", "anomalies"]
    for ws in wb:
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    ws = wb["releves"]
    entetes = [c.value for c in ws[1]]
    col = entetes.index("fichier_photo") + 1
    cell = ws.cell(row=2, column=col)
    assert cell.hyperlink is not None and cell.hyperlink.target == cell.value
    assert (sortie / cell.value).is_file()
    assert ws.max_row == len(res.lignes) + 1
    assert wb["anomalies"].max_row == len(res.anomalies) + 1
    statuts = [r[2] for r in wb["visites"].iter_rows(min_row=2, values_only=True)]
    assert sorted(statuts) == ["corbeille, exclue", "doublon écarté", "importée", "importée"]


def test_fuseau(jeu, tmp_path):
    res, _ = importer(jeu["dossier"], tmp_path / "s", Config(), fuseau="UTC")
    assert str(ligne(res, "e01")["date_heure_locale"]) == "2026-10-05 08:00:02"
    with pytest.raises(ErreurImport):
        importer(jeu["dossier"], tmp_path / "s2", Config(), fuseau="Mars/Olympus")


# ------------------------------------------------------------------ ZIP et index.json

def test_zip_avec_index(jeu, tmp_path):
    res, _ = importer(jeu["zip"], tmp_path / "s", Config())
    assert res.nb_digests_lus == 2 and len(res.retenues) == 1
    idx = [a for a in res.anomalies if a.code == "INDEX_STATUT"]
    assert {a.gravite for a in idx} == {"attention", "erreur"}
    assert any("lecture impossible" in a.message for a in idx)
    assert Counter(l["type"] for l in res.lignes)["photo"] == 3


def test_index_fichier_absent(jeu, tmp_path):
    import zipfile
    z = tmp_path / "incomplet.zip"
    with zipfile.ZipFile(jeu["zip"]) as src, zipfile.ZipFile(z, "w") as dst:
        dst.writestr("index.json", src.read("index.json"))
    res, _ = importer(z, tmp_path / "s", Config())
    assert sum(a.code == "INDEX_FICHIER_ABSENT" for a in res.anomalies) == 2


# ------------------------------------------------------------------ idempotence et sécurité

def test_idempotent_et_entree_intacte(jeu, tmp_path):
    avant = empreintes(jeu["dossier"])
    s = tmp_path / "s"
    importer(jeu["dossier"], s, Config())
    (s / "photos" / "reste-d-un-ancien-import.jpg").write_bytes(b"x")
    premier = {k: v for k, v in empreintes(s).items() if k not in ("releves.xlsx", "rapport_import.txt")}
    importer(jeu["dossier"], s, Config())
    second = {k: v for k, v in empreintes(s).items() if k not in ("releves.xlsx", "rapport_import.txt")}
    del premier["photos/reste-d-un-ancien-import.jpg"]
    assert premier == second
    assert empreintes(jeu["dossier"]) == avant


def test_refus_dossier_etranger_ou_samples(jeu, tmp_path):
    etranger = tmp_path / "docs"
    etranger.mkdir()
    (etranger / "important.docx").write_bytes(b"x")
    with pytest.raises(ErreurImport):
        importer(jeu["dossier"], etranger, Config())
    assert (etranger / "important.docx").exists()
    with pytest.raises(ErreurImport):
        importer(jeu["dossier"], tmp_path / "samples" / "sortie", Config())
    with pytest.raises(ErreurImport):
        importer(jeu["dossier"], jeu["dossier"] / "sortie", Config())
    assert not (tmp_path / "samples").exists() and not (jeu["dossier"] / "sortie").exists()


def test_cli(jeu, tmp_path, capsys):
    assert main(["importer", str(jeu["dossier"]), "--sortie", str(tmp_path / "s"), "--fuseau", "Africa/Kigali"]) == 0
    assert "visite(s) importée(s)" in capsys.readouterr().out
    assert main(["importer", str(tmp_path / "absent"), "--sortie", str(tmp_path / "s2")]) == 2
